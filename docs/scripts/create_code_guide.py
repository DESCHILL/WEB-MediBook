from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = PROJECT_ROOT / 'docs' / 'reports' / 'source' / 'medibook_code_guide.docx'


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9'):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for name, value in {'top': top, 'start': start, 'bottom': bottom, 'end': end}.items():
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    properties.append(header)


def prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    properties.append(cant_split)


def apply_font(run, name='Aptos', size=11, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def format_paragraph(paragraph, space_before=0, space_after=6, line_spacing=1.16):
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def remove_paragraph_borders(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn('w:pBdr'))
    if borders is not None:
        properties.remove(borders)


def add_paragraph(document, text='', style=None, bold_prefix=None):
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        apply_font(first, bold=True)
        remaining = paragraph.add_run(text[len(bold_prefix):])
        apply_font(remaining)
    else:
        run = paragraph.add_run(text)
        apply_font(run)
    format_paragraph(paragraph)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    run = paragraph.add_run(text)
    apply_font(run)
    format_paragraph(paragraph, space_after=2)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style='List Number')
    run = paragraph.add_run(text)
    apply_font(run)
    format_paragraph(paragraph, space_after=3)
    return paragraph


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.04
    paragraph.paragraph_format.keep_together = True
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F3F5F7')
    properties.append(shading)
    for line_index, line in enumerate(text.splitlines()):
        run = paragraph.add_run(line)
        apply_font(run, name='Consolas', size=9)
        if line_index < len(text.splitlines()) - 1:
            run.add_break()
    return paragraph


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = ''
        run = cell.paragraphs[0].add_run(header)
        apply_font(run, size=10, bold=True, color=(255, 255, 255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '1F4E78')
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[index])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = ''
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            apply_font(run, size=9.5)
            format_paragraph(paragraph, space_after=0, line_spacing=1.08)
            if row_index % 2 == 1:
                set_cell_shading(cell, 'F6F8FA')
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Cm(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f'Heading {level}')
    run = paragraph.add_run(text)
    apply_font(run, size=15 if level == 1 else 12, bold=True, color=(0, 0, 0))
    paragraph.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    styles = document.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(11)
    for style_name in ('Title', 'Heading 1', 'Heading 2'):
        style = styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = styles['Title']._element.get_or_add_pPr()
    title_borders = title_properties.find(qn('w:pBdr'))
    if title_borders is not None:
        title_properties.remove(title_borders)
    if 'Code Block' not in styles:
        styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)


def create_document():
    from code_guide_content import write_content
    document = Document()
    configure_document(document)
    title = document.add_paragraph('Hướng dẫn chạy và đọc code DoctorSewa', style='Title')
    for run in title.runs:
        apply_font(run, size=22, bold=True, color=(0, 0, 0))
    remove_paragraph_borders(title)
    write_content(document, add_heading, add_paragraph, add_code, add_table)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == '__main__':
    create_document()
